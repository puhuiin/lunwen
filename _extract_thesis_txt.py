import docx
from pathlib import Path

p = Path("reports/论文初稿_2026-08-26.docx")
d = docx.Document(p)
lines = []
for para in d.paragraphs:
    t = para.text.strip()
    if t:
        lines.append(t)

ti = 0
for t in d.tables:
    ti += 1
    lines.append(f"\n===== TABLE {ti} =====")
    for row in t.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        lines.append(" | ".join(cells))

out = Path("output/_thesis_fulltext.txt")
txt = "\n".join(lines)
out.write_text(txt, encoding="utf-8")
print("total chars:", len(txt))
print("non-empty blocks:", len(lines))
print("tables:", ti)
