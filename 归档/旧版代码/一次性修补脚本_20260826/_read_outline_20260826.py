import docx
from pathlib import Path

P = Path(r'd:\Desktop\基金经理行为分析研究\投资经理画像分析大纲.docx')
d = docx.Document(str(P))
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    print('%3d [%s] %s' % (i, p.style.name, t))
print('--- tables:', len(d.tables))
for ti, tb in enumerate(d.tables):
    print('== table', ti, len(tb.rows), 'x', len(tb.columns))
    for r in tb.rows[:30]:
        print(' | '.join(c.text.strip().replace('\n', ' ') for c in r.cells))
